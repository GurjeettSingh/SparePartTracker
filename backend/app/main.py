from __future__ import annotations

import json
import re
from datetime import datetime
from io import BytesIO

from fastapi import Body, Depends, FastAPI, File, HTTPException, Query, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import func, select, text, tuple_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from . import models
from .config import CORS_ORIGINS
from .db import Base, engine, get_db, get_db_session
from .exporting import build_order_excel, build_order_pdf, safe_export_filename
from .schemas import (
    AuthOut,
    LoginIn,
    ManufacturerOut,
    ModelOut,
    SignupIn,
    OrderCreateWithItems,
    OrderItemCreate,
    OrderItemOut,
    OrderItemUpdate,
    OrderOut,
    OrderRename,
    OrderSummaryOut,
    OrderWithItemsOut,
    SparePartOut,
    UserOut,
    UserProfileUpdate,
    OrderPurchaseUpdate,
    InventoryOut,
    InventoryCreate,
    InventoryUpdate,
    InventoryUpsert,
)


_ORDER_NAME_DATE_SUFFIX_RE = re.compile(r"\s-\s\d{2}-\d{2}-\d{4}$")


def _with_today_date_suffix(order_name: str | None) -> str:
    base = (order_name or "").strip()
    if not base:
        base = "Order"
    # Avoid duplicating an existing DD-MM-YYYY suffix.
    if _ORDER_NAME_DATE_SUFFIX_RE.search(base):
        return base
    today = datetime.now().strftime("%d-%m-%Y")
    return f"{base} - {today}"
from .seed import seed_if_empty
from .auth import create_access_token, hash_password, verify_password
from .deps import get_current_user
from .routers.inventory import router as inventory_router

app = FastAPI(title="Spare Part Tracker API")

# Inventory endpoints (router) — final paths: /inventory and /inventory/{id}
app.include_router(inventory_router, prefix="/inventory", tags=["inventory"])

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def _ensure_schema() -> None:
    # Lightweight migration so existing volumes don't break.
    with engine.begin() as conn:
        conn.execute(
            text(
                """
                ALTER TABLE orders
                ADD COLUMN IF NOT EXISTS order_name VARCHAR(200);
                """
            )
        )
        conn.execute(
            text(
                """
                ALTER TABLE orders
                ADD COLUMN IF NOT EXISTS user_id INTEGER;
                """
            )
        )

        conn.execute(
            text(
                """
                ALTER TABLE orders
                ADD COLUMN IF NOT EXISTS status VARCHAR(20) NOT NULL DEFAULT 'Draft';
                """
            )
        )
        conn.execute(
            text(
                """
                ALTER TABLE orders
                ADD COLUMN IF NOT EXISTS supplier_name VARCHAR(200);
                """
            )
        )

        conn.execute(
            text(
                """
                ALTER TABLE order_items
                ADD COLUMN IF NOT EXISTS available_stock INTEGER NOT NULL DEFAULT 0;
                """
            )
        )
        conn.execute(
            text(
                """
                ALTER TABLE order_items
                ADD COLUMN IF NOT EXISTS to_purchase INTEGER NOT NULL DEFAULT 0;
                """
            )
        )

        # Backfill existing rows (safe to run repeatedly)
        conn.execute(text("UPDATE order_items SET available_stock = 0 WHERE available_stock IS NULL"))
        conn.execute(
            text(
                """
                UPDATE order_items
                SET to_purchase = GREATEST(0, quantity - available_stock);
                """
            )
        )

        conn.execute(
            text(
                """
                ALTER TABLE spare_parts
                ADD COLUMN IF NOT EXISTS category VARCHAR(40) NOT NULL DEFAULT 'Others';
                """
            )
        )

        conn.execute(
            text(
                """
                DO $$
                BEGIN
                    IF EXISTS (
                        SELECT 1 FROM information_schema.tables
                        WHERE table_schema = 'public' AND table_name = 'inventory'
                    ) AND NOT EXISTS (
                        SELECT 1 FROM information_schema.columns
                        WHERE table_schema = 'public' AND table_name = 'inventory' AND column_name = 'manufacturer_id'
                    ) THEN
                        ALTER TABLE inventory RENAME TO inventory_legacy;
                    END IF;
                END$$;

                CREATE TABLE IF NOT EXISTS inventory (
                    id SERIAL PRIMARY KEY,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    manufacturer_id INTEGER NOT NULL REFERENCES manufacturers(id) ON DELETE CASCADE,
                    model_id INTEGER NOT NULL REFERENCES models(id) ON DELETE CASCADE,
                    spare_part_id INTEGER NOT NULL REFERENCES spare_parts(id) ON DELETE CASCADE,
                    stock_quantity INTEGER NOT NULL DEFAULT 0,
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    CONSTRAINT uq_inventory_user_mfr_model_part UNIQUE (user_id, manufacturer_id, model_id, spare_part_id)
                );
                """
            )
        )


def _calc_to_purchase(quantity: int, available_stock: int) -> int:
    return max(0, int(quantity) - int(available_stock))


@app.post("/auth/signup", response_model=AuthOut)
def signup(payload: SignupIn, db: Session = Depends(get_db)):
    if payload.password != payload.confirm_password:
        raise HTTPException(status_code=400, detail="Password and confirm password do not match")

    existing = (
        db.execute(select(models.User).where(models.User.mobile_number == payload.mobile_number))
        .scalars()
        .first()
    )
    if existing is not None:
        raise HTTPException(status_code=409, detail="User already exists")

    user = models.User(
        first_name=payload.first_name,
        last_name=payload.last_name,
        workshop_name=payload.workshop_name,
        mobile_number=payload.mobile_number,
        email=payload.email,
        password_hash=hash_password(payload.password),
    )
    db.add(user)
    db.commit()
    db.refresh(user)

    token = create_access_token(user_id=user.id, mobile_number=user.mobile_number)
    return AuthOut(
        token=token,
        user=UserOut(
            id=user.id,
            first_name=user.first_name,
            last_name=user.last_name,
            workshop_name=user.workshop_name,
            mobile_number=user.mobile_number,
            email=user.email,
            created_at=user.created_at,
        ),
    )


@app.post("/auth/login", response_model=AuthOut)
def login(payload: LoginIn, db: Session = Depends(get_db)):
    user = (
        db.execute(select(models.User).where(models.User.mobile_number == payload.mobile_number))
        .scalars()
        .first()
    )
    if user is None or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Invalid mobile number or password")

    token = create_access_token(user_id=user.id, mobile_number=user.mobile_number)
    return AuthOut(
        token=token,
        user=UserOut(
            id=user.id,
            first_name=user.first_name,
            last_name=user.last_name,
            workshop_name=user.workshop_name,
            mobile_number=user.mobile_number,
            email=user.email,
            created_at=user.created_at,
        ),
    )


@app.get("/user/profile", response_model=UserOut)
def get_profile(current_user: models.User = Depends(get_current_user)):
    return UserOut(
        id=current_user.id,
        first_name=current_user.first_name,
        last_name=current_user.last_name,
        workshop_name=current_user.workshop_name,
        mobile_number=current_user.mobile_number,
        email=current_user.email,
        created_at=current_user.created_at,
    )


@app.put("/user/profile", response_model=UserOut)
def update_profile(
    payload: UserProfileUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    current_user.workshop_name = payload.workshop_name
    current_user.email = payload.email
    db.add(current_user)
    db.commit()
    db.refresh(current_user)
    return UserOut(
        id=current_user.id,
        first_name=current_user.first_name,
        last_name=current_user.last_name,
        workshop_name=current_user.workshop_name,
        mobile_number=current_user.mobile_number,
        email=current_user.email,
        created_at=current_user.created_at,
    )


@app.on_event("startup")
def on_startup() -> None:
    _ensure_schema()
    Base.metadata.create_all(bind=engine)
    db = get_db_session()
    try:
        seed_if_empty(db)
    finally:
        db.close()


@app.get("/manufacturers", response_model=list[ManufacturerOut])
def list_manufacturers(db: Session = Depends(get_db)):
    rows = db.execute(select(models.Manufacturer).order_by(models.Manufacturer.name)).scalars().all()
    return rows


@app.get("/models", response_model=list[ModelOut])
def list_models(
    manufacturer_id: int = Query(..., ge=1),
    db: Session = Depends(get_db),
):
    rows = (
        db.execute(
            select(models.CarModel)
            .where(models.CarModel.manufacturer_id == manufacturer_id)
            .order_by(models.CarModel.name)
        )
        .scalars()
        .all()
    )
    return rows


@app.get("/spare-parts", response_model=list[SparePartOut])
def list_spare_parts(db: Session = Depends(get_db)):
    rows = db.execute(select(models.SparePart).order_by(models.SparePart.name)).scalars().all()
    return rows


@app.post("/order", response_model=OrderOut)
def create_order(
    payload: OrderCreateWithItems | None = Body(default=None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    # If no payload, create a draft order (unsaved).
    if payload is None:
        order = models.Order(order_name=None, user_id=current_user.id, status="Draft")
        db.add(order)
        db.commit()
        db.refresh(order)
        return OrderOut(
            id=order.id,
            order_name=order.order_name,
            status=order.status,
            supplier_name=order.supplier_name,
            created_at=order.created_at,
        )

    order = models.Order(order_name=_with_today_date_suffix(payload.order_name), user_id=current_user.id, status="Draft")
    db.add(order)
    db.flush()

    # Merge duplicates within the incoming payload
    merged: dict[tuple[int, int], dict[str, int]] = {}
    for it in payload.items:
        key = (it.model_id, it.spare_part_id)
        if key not in merged:
            merged[key] = {"quantity": 0, "available_stock": 0}
        merged[key]["quantity"] += it.quantity
        merged[key]["available_stock"] += max(0, int(it.available_stock or 0))

    for (model_id, spare_part_id), agg in merged.items():
        model = db.get(models.CarModel, model_id)
        if model is None:
            raise HTTPException(status_code=400, detail=f"Invalid model_id: {model_id}")
        manufacturer_id = model.manufacturer_id
        if db.get(models.SparePart, spare_part_id) is None:
            raise HTTPException(status_code=400, detail=f"Invalid spare_part_id: {spare_part_id}")

        qty = int(agg["quantity"])
        stock = int(agg["available_stock"])

        db.add(
            models.OrderItem(
                order_id=order.id,
                manufacturer_id=manufacturer_id,
                model_id=model_id,
                spare_part_id=spare_part_id,
                quantity=qty,
                available_stock=stock,
                to_purchase=_calc_to_purchase(qty, stock),
            )
        )

    db.commit()
    db.refresh(order)
    return OrderOut(
        id=order.id,
        order_name=order.order_name,
        status=order.status,
        supplier_name=order.supplier_name,
        created_at=order.created_at,
    )


@app.get("/orders", response_model=list[OrderSummaryOut])
def list_orders(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    rows = (
        db.execute(
            select(
                models.Order.id,
                models.Order.order_name,
                models.Order.created_at,
                models.Order.status,
                models.Order.supplier_name,
                func.count(models.OrderItem.id).label("total_items"),
            )
            .join(models.OrderItem, models.OrderItem.order_id == models.Order.id, isouter=True)
            .where(models.Order.order_name.is_not(None))
            .where(models.Order.user_id == current_user.id)
            .group_by(models.Order.id)
            .order_by(models.Order.created_at.desc())
        )
        .all()
    )
    return [
        OrderSummaryOut(
            id=r.id,
            order_name=r.order_name,
            created_at=r.created_at,
            total_items=int(r.total_items or 0),
            status=r.status or "Draft",
            supplier_name=r.supplier_name,
        )
        for r in rows
        if r.order_name is not None
    ]


@app.get("/orders/latest", response_model=OrderWithItemsOut)
def get_latest_order(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    order = (
        db.execute(
            select(models.Order)
            .where(models.Order.user_id == current_user.id)
            .where(models.Order.order_name.is_not(None))
            .order_by(models.Order.created_at.desc(), models.Order.id.desc())
            .limit(1)
        )
        .scalars()
        .first()
    )
    if order is None:
        raise HTTPException(status_code=404, detail="No previous order")

    items = db.execute(select(models.OrderItem).where(models.OrderItem.order_id == order.id)).scalars().all()
    return _items_to_schema(db, order, items)


@app.patch("/order/{order_id}", response_model=OrderOut)
def rename_order(
    order_id: int,
    payload: OrderRename,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    order = db.get(models.Order, order_id)
    if order is None or order.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Order not found")

    order.order_name = payload.order_name
    db.add(order)
    db.commit()
    db.refresh(order)
    return OrderOut(
        id=order.id,
        order_name=order.order_name,
        status=order.status,
        supplier_name=order.supplier_name,
        created_at=order.created_at,
    )


@app.patch("/order/{order_id}/purchase", response_model=OrderOut)
def update_purchase_status(
    order_id: int,
    payload: OrderPurchaseUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    order = db.get(models.Order, order_id)
    if order is None or order.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Order not found")

    order.status = payload.status
    order.supplier_name = payload.supplier_name
    db.add(order)
    db.commit()
    db.refresh(order)
    return OrderOut(
        id=order.id,
        order_name=order.order_name,
        status=order.status,
        supplier_name=order.supplier_name,
        created_at=order.created_at,
    )


@app.delete("/order/{order_id}", status_code=204)
def delete_order(order_id: int, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    order = db.get(models.Order, order_id)
    if order is None or order.user_id != current_user.id:
        return Response(status_code=204)
    db.delete(order)
    db.commit()
    return Response(status_code=204)


def _items_to_schema(db: Session, order: models.Order, items: list[models.OrderItem]) -> OrderWithItemsOut:
    for item in items:
        item.manufacturer  # ensure loaded
        item.model
        item.spare_part

    keys = {(it.manufacturer_id, it.model_id, it.spare_part_id) for it in items}
    catalog_by_key: dict[tuple[int, int, int], models.PartCatalog] = {}
    if keys:
        rows = (
            db.execute(
                select(models.PartCatalog).where(
                    tuple_(
                        models.PartCatalog.manufacturer_id,
                        models.PartCatalog.model_id,
                        models.PartCatalog.spare_part_id,
                    ).in_(keys)
                )
            )
            .scalars()
            .all()
        )
        catalog_by_key = {
            (r.manufacturer_id, r.model_id, r.spare_part_id): r
            for r in rows
        }

    return OrderWithItemsOut(
        id=order.id,
        order_name=order.order_name,
        status=order.status,
        supplier_name=order.supplier_name,
        created_at=order.created_at,
        items=[
            OrderItemOut(
                id=item.id,
                order_id=item.order_id,
                manufacturer_id=item.manufacturer_id,
                manufacturer_name=item.manufacturer.name,
                model_id=item.model_id,
                model_name=item.model.name,
                spare_part_id=item.spare_part_id,
                spare_part_name=item.spare_part.name,
                spare_part_category=getattr(item.spare_part, "category", "Others") or "Others",
                quantity=item.quantity,
                available_stock=getattr(item, "available_stock", 0) or 0,
                to_purchase=getattr(item, "to_purchase", 0) or _calc_to_purchase(item.quantity, getattr(item, "available_stock", 0) or 0),
                typical_specification=catalog_by_key.get(
                    (item.manufacturer_id, item.model_id, item.spare_part_id)
                ).typical_specification
                if catalog_by_key.get((item.manufacturer_id, item.model_id, item.spare_part_id))
                else None,
                oem_part_number=catalog_by_key.get(
                    (item.manufacturer_id, item.model_id, item.spare_part_id)
                ).oem_part_number
                if catalog_by_key.get((item.manufacturer_id, item.model_id, item.spare_part_id))
                else None,
            )
            for item in items
        ],
    )


@app.get("/order/{order_id}", response_model=OrderWithItemsOut)
def get_order(order_id: int, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    order = db.get(models.Order, order_id)
    if order is None or order.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Order not found")

    # Load items with names
    items = db.execute(select(models.OrderItem).where(models.OrderItem.order_id == order_id)).scalars().all()
    return _items_to_schema(db, order, items)


@app.post("/order-item", response_model=OrderItemOut)
def add_order_item(payload: OrderItemCreate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    order = db.get(models.Order, payload.order_id)
    if order is None or order.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Order not found")

    model = db.get(models.CarModel, payload.model_id)
    if model is None:
        raise HTTPException(status_code=400, detail="Invalid model")
    if model.manufacturer_id != payload.manufacturer_id:
        raise HTTPException(status_code=400, detail="Model does not belong to manufacturer")

    if db.get(models.Manufacturer, payload.manufacturer_id) is None:
        raise HTTPException(status_code=400, detail="Invalid manufacturer")
    if db.get(models.SparePart, payload.spare_part_id) is None:
        raise HTTPException(status_code=400, detail="Invalid spare part")

    item = models.OrderItem(
        order_id=payload.order_id,
        manufacturer_id=payload.manufacturer_id,
        model_id=payload.model_id,
        spare_part_id=payload.spare_part_id,
        quantity=payload.quantity,
        available_stock=payload.available_stock,
        to_purchase=_calc_to_purchase(payload.quantity, payload.available_stock),
    )
    db.add(item)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        existing = (
            db.execute(
                select(models.OrderItem)
                .where(models.OrderItem.order_id == payload.order_id)
                .where(models.OrderItem.model_id == payload.model_id)
                .where(models.OrderItem.spare_part_id == payload.spare_part_id)
            )
            .scalars()
            .first()
        )
        if existing is None:
            raise HTTPException(status_code=409, detail="Duplicate item")

        existing.quantity += payload.quantity
        existing.to_purchase = _calc_to_purchase(existing.quantity, getattr(existing, "available_stock", 0) or 0)
        db.add(existing)
        db.commit()
        db.refresh(existing)
        existing.manufacturer
        existing.model
        existing.spare_part

        return OrderItemOut(
            id=existing.id,
            order_id=existing.order_id,
            manufacturer_id=existing.manufacturer_id,
            manufacturer_name=existing.manufacturer.name,
            model_id=existing.model_id,
            model_name=existing.model.name,
            spare_part_id=existing.spare_part_id,
            spare_part_name=existing.spare_part.name,
            quantity=existing.quantity,
            spare_part_category=getattr(existing.spare_part, "category", "Others") or "Others",
            available_stock=getattr(existing, "available_stock", 0) or 0,
            to_purchase=getattr(existing, "to_purchase", 0) or _calc_to_purchase(existing.quantity, getattr(existing, "available_stock", 0) or 0),
        )

    db.refresh(item)
    item.manufacturer
    item.model
    item.spare_part

    return OrderItemOut(
        id=item.id,
        order_id=item.order_id,
        manufacturer_id=item.manufacturer_id,
        manufacturer_name=item.manufacturer.name,
        model_id=item.model_id,
        model_name=item.model.name,
        spare_part_id=item.spare_part_id,
        spare_part_name=item.spare_part.name,
        quantity=item.quantity,
        spare_part_category=getattr(item.spare_part, "category", "Others") or "Others",
        available_stock=getattr(item, "available_stock", 0) or 0,
        to_purchase=getattr(item, "to_purchase", 0) or _calc_to_purchase(item.quantity, getattr(item, "available_stock", 0) or 0),
    )


@app.patch("/order-item/{item_id}", response_model=OrderItemOut)
def update_order_item(item_id: int, payload: OrderItemUpdate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    item = db.get(models.OrderItem, item_id)
    if item is None:
        raise HTTPException(status_code=404, detail="Item not found")

    order = db.get(models.Order, item.order_id)
    if order is None or order.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Item not found")

    if payload.quantity is None and payload.available_stock is None:
        raise HTTPException(status_code=400, detail="Provide quantity and/or available_stock")

    if payload.quantity is not None:
        item.quantity = payload.quantity
    if payload.available_stock is not None:
        item.available_stock = payload.available_stock
    item.to_purchase = _calc_to_purchase(item.quantity, getattr(item, "available_stock", 0) or 0)
    db.add(item)
    db.commit()
    db.refresh(item)

    item.manufacturer
    item.model
    item.spare_part

    return OrderItemOut(
        id=item.id,
        order_id=item.order_id,
        manufacturer_id=item.manufacturer_id,
        manufacturer_name=item.manufacturer.name,
        model_id=item.model_id,
        model_name=item.model.name,
        spare_part_id=item.spare_part_id,
        spare_part_name=item.spare_part.name,
        quantity=item.quantity,
        spare_part_category=getattr(item.spare_part, "category", "Others") or "Others",
        available_stock=getattr(item, "available_stock", 0) or 0,
        to_purchase=getattr(item, "to_purchase", 0) or _calc_to_purchase(item.quantity, getattr(item, "available_stock", 0) or 0),
    )


@app.delete("/order-item/{item_id}", status_code=204)
def delete_order_item(item_id: int, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    item = db.get(models.OrderItem, item_id)
    if item is None:
        return Response(status_code=204)

    order = db.get(models.Order, item.order_id)
    if order is None or order.user_id != current_user.id:
        return Response(status_code=204)

    db.delete(item)
    db.commit()
    return Response(status_code=204)


@app.get("/export/pdf/{order_id}")
def export_pdf(order_id: int, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    order = db.get(models.Order, order_id)
    if order is None or order.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Order not found")

    items = db.execute(select(models.OrderItem).where(models.OrderItem.order_id == order_id)).scalars().all()
    order_schema = _items_to_schema(db, order, items)
    pdf_bytes = build_order_pdf(order_schema, workshop_name=getattr(current_user, "workshop_name", None))
    filename = safe_export_filename(order_schema.order_name or f"Order {order_id}", "pdf")
    return StreamingResponse(
        iter([pdf_bytes]),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=\"{filename}\""},
    )


@app.get("/export/excel/{order_id}")
def export_excel(order_id: int, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    order = db.get(models.Order, order_id)
    if order is None or order.user_id != current_user.id:
        raise HTTPException(status_code=404, detail="Order not found")

    items = db.execute(select(models.OrderItem).where(models.OrderItem.order_id == order_id)).scalars().all()
    order_schema = _items_to_schema(db, order, items)
    xlsx_bytes = build_order_excel(order_schema, workshop_name=getattr(current_user, "workshop_name", None))
    filename = safe_export_filename(order_schema.order_name or f"Order {order_id}", "xlsx")
    return StreamingResponse(
        iter([xlsx_bytes]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{filename}\""},
    )


def _norm(s: str) -> str:
    return " ".join(s.strip().lower().split())


def _norm_header(s: str) -> str:
    # For header matching: treat underscores/dashes as spaces.
    return _norm(s.replace("_", " ").replace("-", " "))


@app.post("/import-order")
async def import_order(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    raw = await file.read()
    filename = file.filename or ""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    rows: list[dict[str, str]] = []

    if ext == "json":
        try:
            data = json.loads(raw.decode("utf-8"))
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid JSON file")

        items = data.get("items") if isinstance(data, dict) else data
        if not isinstance(items, list):
            raise HTTPException(status_code=400, detail="JSON must be an array or {items:[...]}" )

        for it in items:
            if not isinstance(it, dict):
                continue
            rows.append(
                {
                    "manufacturer": str(it.get("manufacturer") or it.get("manufacturer_name") or ""),
                    "model": str(it.get("model") or it.get("model_name") or ""),
                    "spare_part": str(it.get("spare_part") or it.get("spare_part_name") or ""),
                    "quantity": str(it.get("quantity") or ""),
                }
            )
    elif ext == "xlsx":
        try:
            wb = load_workbook(filename=BytesIO(raw), data_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid Excel file")

        ws = wb.active
        values = list(ws.values)
        if not values:
            raise HTTPException(status_code=400, detail="Excel file is empty")

        required = {
            "manufacturer": "Manufacturer",
            "model": "Model",
            "spare part": "Spare Part",
            "quantity": "Quantity",
        }

        # Find the header row (supports files that start with title/date rows).
        header_row_idx: int | None = None
        header_idx: dict[str, int] = {}
        found_headers: list[str] = []

        for ridx, row in enumerate(values[:25]):
            if not row:
                continue
            hdr = [str(c or "") for c in row]
            idx: dict[str, int] = {}
            for i, name in enumerate(hdr):
                key = _norm_header(name)
                if key:
                    idx[key] = i

            # Accept common export header alias.
            if "quantity" not in idx and "required quantity" in idx:
                idx["quantity"] = idx["required quantity"]

            missing = [k for k in required.keys() if k not in idx]
            if not missing:
                header_row_idx = ridx
                header_idx = idx
                found_headers = hdr
                break

        if header_row_idx is None:
            # Best-effort list of headers from the first non-empty row.
            for row in values[:10]:
                if row:
                    found_headers = [str(c or "") for c in row]
                    break

            raise HTTPException(
                status_code=400,
                detail=(
                    "Excel must have columns: Manufacturer, Model, Spare Part, Quantity. "
                    + (f"Found headers: {found_headers}" if found_headers else "No header row found")
                ),
            )

        mfr_col = header_idx["manufacturer"]
        model_col = header_idx["model"]
        part_col = header_idx["spare part"]
        qty_col = header_idx["quantity"]

        for excel_row_index, row in enumerate(values[header_row_idx + 1 :], start=header_row_idx + 2):
            if row is None:
                continue

            def val(i: int) -> str:
                return str(row[i] or "").strip() if i < len(row) else ""

            # skip empty rows
            if not (val(mfr_col) or val(model_col) or val(part_col) or val(qty_col)):
                continue

            rows.append(
                {
                    "manufacturer": val(mfr_col),
                    "model": val(model_col),
                    "spare_part": val(part_col),
                    "quantity": val(qty_col),
                    "_excel_row": str(excel_row_index),
                }
            )
    else:
        raise HTTPException(status_code=400, detail="Upload a .xlsx or .json file")

    if not rows:
        raise HTTPException(status_code=400, detail="No rows found")

    manufacturers = db.execute(select(models.Manufacturer)).scalars().all()
    mfr_by_norm = {_norm(m.name): m for m in manufacturers}

    parts = db.execute(select(models.SparePart)).scalars().all()
    part_by_norm = {_norm(p.name): p for p in parts}

    all_models = db.execute(select(models.CarModel)).scalars().all()
    model_by_norm: dict[tuple[int, str], models.CarModel] = {}
    for m in all_models:
        model_by_norm[(m.manufacturer_id, _norm(m.name))] = m

    errors: list[str] = []
    imported: list[dict[str, object]] = []

    for row_num, r in enumerate(rows, start=2 if ext == "xlsx" else 1):
        if ext == "xlsx":
            try:
                row_num = int(r.get("_excel_row", str(row_num)))
            except Exception:
                pass

        mfr_name = r.get("manufacturer", "").strip()
        model_name = r.get("model", "").strip()
        part_name = r.get("spare_part", "").strip()
        qty_raw = r.get("quantity", "").strip()

        if not mfr_name or not model_name or not part_name or not qty_raw:
            errors.append(f"Row {row_num}: Missing required fields")
            continue

        try:
            qty = int(float(qty_raw))
        except Exception:
            errors.append(f"Row {row_num}: Quantity must be a number")
            continue
        if qty <= 0:
            errors.append(f"Row {row_num}: Quantity must be > 0")
            continue

        mfr = mfr_by_norm.get(_norm(mfr_name))
        if mfr is None:
            errors.append(f"Row {row_num}: Unknown manufacturer '{mfr_name}'")
            continue

        model = model_by_norm.get((mfr.id, _norm(model_name)))
        if model is None:
            errors.append(f"Row {row_num}: Unknown model '{model_name}' for '{mfr.name}'")
            continue

        part = part_by_norm.get(_norm(part_name))
        if part is None:
            errors.append(f"Row {row_num}: Unknown spare part '{part_name}'")
            continue

        imported.append(
            {
                "manufacturer_id": mfr.id,
                "manufacturer_name": mfr.name,
                "model_id": model.id,
                "model_name": model.name,
                "spare_part_id": part.id,
                "spare_part_name": part.name,
                "quantity": qty,
                "available_stock": 0,
            }
        )

    if errors:
        raise HTTPException(status_code=400, detail=errors)

    return {"items": imported}
